"""Generate an editable PulseSplit report in the style of the supplied Horizon sample."""
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_COLOR_INDEX

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'docs';FIG=ROOT/'figures'
BLUE=RGBColor(11,83,148)

def shade(cell,fill='D9E8F5'):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd');shd.set(qn('w:fill'),fill);tcPr.append(shd)
def set_cell_border(cell):
    tcPr=cell._tc.get_or_add_tcPr();bd=tcPr.first_child_found_in('w:tcBorders')
    if bd is None:bd=OxmlElement('w:tcBorders');tcPr.append(bd)
    for s in ('top','left','bottom','right'):
        elem=OxmlElement('w:'+s);elem.set(qn('w:val'),'single');elem.set(qn('w:sz'),'4');elem.set(qn('w:color'),'777777');bd.append(elem)

def field(page):
    begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin')
    t=OxmlElement('w:instrText');t.set(qn('xml:space'),'preserve');t.text=' PAGE '
    sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate')
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end')
    r=page.add_run();r._r.append(begin);r=page.add_run();r._r.append(t);r=page.add_run();r._r.append(sep);r=page.add_run();r._r.append(end)

def page_format(section,fmt='lowerRoman',start=1):
    secPr=section._sectPr
    p=secPr.find(qn('w:pgNumType'))
    if p is None:p=OxmlElement('w:pgNumType');secPr.append(p)
    p.set(qn('w:fmt'),fmt);p.set(qn('w:start'),str(start))

doc=Document();sect=doc.sections[0]
sect.page_height=Inches(11.694);sect.page_width=Inches(8.268)
sect.left_margin=Inches(1.50);sect.right_margin=Inches(.75)
sect.top_margin=Inches(.90);sect.bottom_margin=Inches(.85)
sect.header_distance=Inches(.35);sect.footer_distance=Inches(.42)
page_format(sect)
base=doc.styles['Normal'];base.font.name='Times New Roman';base.font.size=Pt(12);base.font.color.rgb=RGBColor(0,0,0)
base.paragraph_format.line_spacing=1.55;base.paragraph_format.space_after=Pt(9)
base.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name,size in [('Report Chapter',16),('Report Subheading',12),('Prelim Heading',16)]:
    style=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH);style.font.name='Times New Roman';style.font.size=Pt(size);style.font.bold=True;style.font.color.rgb=BLUE
    style.paragraph_format.space_after=Pt(16 if size==16 else 10);style.paragraph_format.keep_with_next=True
    style.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER if size==16 else WD_ALIGN_PARAGRAPH.LEFT
for name in ['Caption','Caption Text']:
    s=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    s.font.name='Times New Roman';s.font.size=Pt(11);s.font.italic=True
    s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;s.paragraph_format.space_before=Pt(5);s.paragraph_format.space_after=Pt(12)
footer=sect.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;field(footer)

def page():doc.add_page_break()
def para(text='',bold_intro=None):
    p=doc.add_paragraph()
    if bold_intro:
        p.add_run(bold_intro).bold=True;p.add_run(text)
    else:p.add_run(text)
    return p

def heading(text):doc.add_paragraph(text,style='Report Subheading')
def chapter(text):
    p=doc.add_paragraph('Chapter '+text,style='Report Chapter')
    p.paragraph_format.page_break_before=True
def prelim(text):
    page();doc.add_paragraph(text,style='Prelim Heading')
def bullet(items):
    for item in items:
        p=doc.add_paragraph(style='List Bullet');p.add_run(item)

def table(headers,rows,caption=None):
    tab=doc.add_table(rows=1,cols=len(headers));tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
    usable=Inches(6.0)
    for i,h in enumerate(headers):
        cell=tab.rows[0].cells[i];cell.text=str(h);shade(cell)
        for p in cell.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:r.bold=True;r.font.name='Times New Roman';r.font.size=Pt(10)
        set_cell_border(cell)
    for row in rows:
        cells=tab.add_row().cells
        for j,item in enumerate(row):
            cells[j].text=str(item);cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[j].paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10)
            set_cell_border(cells[j])
    if caption:doc.add_paragraph(caption,style='Caption')
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return tab

def figure(name,caption):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIG/name),width=Inches(5.9))
    doc.add_paragraph(caption,style='Caption')

# Title page modeled on reference report: large whitespace, red title, blue department.
p=doc.add_paragraph();p.paragraph_format.space_before=Pt(70);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Deep Learning Report');r.font.size=Pt(20)
p=doc.add_paragraph();p.paragraph_format.space_before=Pt(24);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('PULSESPLIT: DEEP LEARNING-BASED\nHEART AND LUNG SOUND SEPARATION');r.bold=True;r.font.size=Pt(16);r.font.color.rgb=RGBColor(255,0,0)
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('CSE Department\n\nIII-Year - I-Semester');r.bold=True;r.font.size=Pt(14);r.font.color.rgb=BLUE
p=doc.add_paragraph();p.paragraph_format.space_before=Pt(27);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('SUBMITTED BY');r.bold=True;r.font.size=Pt(14)
t=doc.add_table(rows=3,cols=2)
for row,(n,urn) in zip(t.rows,[('[MEMBER 1]','[URN 1]'),('[MEMBER 2]','[URN 2]'),('[MEMBER 3]','[URN 3]')]):
    row.cells[0].text=n;row.cells[1].text=urn
    for c in row.cells:
        for pp in c.paragraphs:
            pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for rr in pp.runs:rr.font.size=Pt(14)
p=doc.add_paragraph();p.paragraph_format.space_before=Pt(20);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('UNDER THE GUIDANCE OF\n\n[GUIDE NAME]\n[Designation, Department of CSE]\nSreenidhi University - Hyderabad\n\n2026-2027');r.font.size=Pt(13)

prelim('Acknowledgment')
para('We are grateful for the opportunity to undertake the technical project titled “PulseSplit: Deep Learning-Based Heart and Lung Sound Separation” as part of our Third Year, First Semester academic curriculum at Sreenidhi University during 2026–2027.')
para('We thank the Department of Computer Science and Engineering for providing an environment that encourages applied research and reproducible engineering. The names and designations of administrative staff and the academic project guide will be inserted after confirmation, rather than assumed from another team’s report.')
para('We acknowledge Yasaman Torabi, Shahram Shirani and James P. Reilly for making the HLS-CMDS dataset available through the UCI Machine Learning Repository. Their documented collection procedure allows careful study of single-channel cardiopulmonary sound separation.')
para('We also acknowledge the maintainers of PyTorch, NumPy, SciPy, scikit-learn, Matplotlib and Streamlit. We will acknowledge all additional contributors and resources used in the final submitted implementation.')

prelim('Student Declaration')
para('We declare that the project report entitled “PulseSplit: Deep Learning-Based Heart and Lung Sound Separation” will document our own implementation and evaluation. It is intended as a reproducible experimental study for a Bachelor of Technology project in Computer Science and Engineering at Sreenidhi University during 2026–2027.')
para('Any model-performance statements in the final report will be derived from the actual executed experiment with the authorized dataset. No simulated smoke-test results will be represented as genuine model accuracy. The report will cite the original dataset authors and relevant research papers.')
para('This is a draft declaration pending confirmation of contributors, academic guide, completed experiments and individual signatures.')
table(['Contributor','URN','Signature'],[['[MEMBER 1]','[URN 1]',''],['[MEMBER 2]','[URN 2]',''],['[MEMBER 3]','[URN 3]','']])

prelim('CERTIFICATE')
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run('Yamnampet, Ghatkesar, Hyderabad 501301')
para('This certificate page has been prepared in the style of the supplied reference report. The formal attestation must be completed and approved by the authorized guide and department after the project work has been reviewed. It must not be presented as an already signed certificate.')
para('Provisional project title: “PulseSplit: Deep Learning-Based Heart and Lung Sound Separation”')
para('Student contributors: [MEMBER 1], [MEMBER 2], [MEMBER 3]. Academic year: 2026–2027.')
for _ in range(4):para('')
table(['Internal Guide','Head of Department','External Examiner'],[['[GUIDE NAME]','[HOD NAME]','[EXAMINER NAME]'],['Signature: __________','Signature: __________','Signature: __________'],['Date: __________','Date: __________','Date: __________']])

prelim('Abstract')
para('Cardiopulmonary auscultation recordings often contain overlapping heart and lung sounds. Separating the two components from a single recorded channel is difficult because both occupy partially overlapping frequency ranges and their relative intensities vary with anatomical position. Methods that analyze a mixed recording without source separation may conceal clinically relevant acoustic patterns. Deep learning offers an approach to learning spectrotemporal structure directly from paired examples, although accuracy on simulated recordings does not demonstrate clinical utility.')
para('PulseSplit is a proposed supervised deep learning system for separating combined heart-lung recordings into estimated heart-only and lung-only signals. The study uses the HLS-CMDS dataset hosted by the UCI Machine Learning Repository, which provides 145 mixed recordings with paired isolated sources, alongside additional standalone examples. Each recording is sampled at 22,050 Hz and lasts 15 seconds. The implementation includes train-only preprocessing, recording-level partitioning, short-time Fourier transform features and a compact encoder-decoder convolutional neural network.')
para('The U-Net estimates two complementary time-frequency masks and reconstructs output waveforms using the mixture phase. The evaluation plan compares scale-invariant signal-to-distortion ratio on withheld recordings with an unprocessed-mixture baseline, alongside listening and spectrogram examination. The implementation includes code-level tests and an interactive audio playback interface.')
para('At the time of this draft, the real dataset has not been downloaded into the execution environment; therefore, actual test-set scores and performance conclusions are intentionally pending. The work is for research and education only, not diagnosis or clinical deployment.')
para('Keywords: Heart–Lung Sound Separation, Deep Learning, U-Net, Spectrogram, Audio Source Separation, HLS-CMDS, SI-SDR, Digital Stethoscope.')

prelim('List of Figures')
table(['Figure No.','Figure Title','Page No.'],[['3.1','PulseSplit System Architecture','5'],['3.2','Spectrogram-Mask U-Net Layer Configuration','6'],['3.3','Experiment and Evaluation Workflow','7']])
prelim('List of Tables')
table(['Table No.','Table Title','Page No.'],[['2.1','Review of Relevant Research','4'],['3.1','Functional and Non-functional Requirements','5–6'],['4.1','Dataset Composition and Roles','8'],['4.2','Model and Training Parameters','9–10'],['6.1','Held-out Evaluation Table','13']])
prelim('List of Abbreviations')
table(['Abbreviation','Full Form'],[['AI','Artificial Intelligence'],['CNN','Convolutional Neural Network'],['DL','Deep Learning'],['FFT','Fast Fourier Transform'],['HLS-CMDS','Heart and Lung Sounds Dataset from a Clinical Manikin using Digital Stethoscope'],['HS','Heart Sound'],['LS','Lung Sound'],['MAE','Mean Absolute Error'],['SI-SDR','Scale-Invariant Signal-to-Distortion Ratio'],['STFT','Short-Time Fourier Transform'],['U-Net','Encoder–Decoder Network with Skip Connections'],['WAV','Waveform Audio File']])
prelim('Table of Contents')
table(['Topic','Page No.'],[['Preliminary Pages','i–ix'],['Chapter 1: Introduction','1'],['1.1 Problem Domain; 1.2 Problem Statement; 1.3 Motivation; 1.4 Objectives; 1.5 Scope','1–2'],['Chapter 2: Literature Survey','3'],['2.1 Review; 2.2 Related Work; 2.3 Comparison; 2.4 Gaps','3–4'],['Chapter 3: System Analysis and Design','5'],['3.1 Existing System; 3.2 Proposed System; 3.3 Requirements; 3.4 Architecture','5–7'],['Chapter 4: Dataset and Methodology','8'],['4.1 Data Source; 4.2 Preparation; 4.3 Model; 4.4 Training; 4.5 Evaluation','8–10'],['Chapter 5: Implementation','11'],['5.1 Environment; 5.2 Modules; 5.3 Demo; 5.4 Testing','11–12'],['Chapter 6: Results and Discussion','13'],['Chapter 7: Conclusion and Future Scope','15'],['References and Appendices','16']])

body=doc.add_section(WD_SECTION_START.NEW_PAGE)
body.page_height=sect.page_height;body.page_width=sect.page_width
body.left_margin=sect.left_margin;body.right_margin=sect.right_margin
body.top_margin=sect.top_margin;body.bottom_margin=sect.bottom_margin
body.footer_distance=sect.footer_distance
body.footer.is_linked_to_previous=False
p=body.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;field(p)
page_format(body,'decimal',1)
# First chapter doesn't need additional page break after section.
doc.add_paragraph('Chapter 1: Introduction',style='Report Chapter')
heading('1.1 Introduction to the Problem Domain')
para('Auscultation provides a way to listen to physiological sound events originating in the cardiovascular and respiratory systems. Heart sounds frequently contain quasi-periodic transient components related to the cardiac cycle, while lung sounds can exhibit broader spectral structure shaped by breathing. A single digital stethoscope often records both sources at once. Their overlap may complicate downstream automated analysis and the interpretation of faint acoustic events.')
para('Audio source separation attempts to decompose an observed signal into constituent signals. Single-channel separation is more challenging than microphone-array separation because only one measured waveform is available, with no spatial cues from multiple microphones. Deep learning methods can learn statistical regularities of different source categories, but they require careful training design and honest testing under conditions that reflect the intended use case.')
para('PulseSplit investigates whether a compact spectrogram-based convolutional encoder-decoder can estimate heart and lung signals from mixed recordings. The project explicitly separates three questions: whether software operates correctly, whether the model fits a controlled audio dataset, and whether such performance would generalize beyond the data source. Only the first is presently tested in this draft.')
heading('1.2 Problem Statement')
para('Given a mono recording x(t) containing simultaneously audible heart and lung sounds, develop a reproducible method that estimates h(t) and l(t) as two output waveforms. The model should capture useful time-frequency patterns from paired examples and preserve the characteristic content of each source as much as possible. Output quality must be examined with quantitative signal metrics and transparent listening examples.')
para('The project is evaluated on paired source recordings made from a clinical manikin. Its objective is source estimation, not identification of cardiac or respiratory disease, and no clinical diagnostic performance is claimed.')
heading('1.3 Motivation')
para('This problem has an unusual combination of signal processing and deep learning: loading and aligning waveforms, transforming them into time-frequency representations, predicting source masks, reconstructing audio, and measuring separation quality. A working demonstration is also immediately understandable because users can hear the mixture and both estimated outputs.')
heading('1.4 Objectives')
bullet(['Use only an approved dataset source: HLS-CMDS from UCI.','Form paired triplets and audit all required audio paths and metadata.','Avoid leakage by partitioning recording IDs before creating waveform segments.','Develop a trainable compact U-Net for two-source separation.','Compare estimates with an explicit unprocessed-mixture baseline using SI-SDR and improvement.','Publish reproducible code, tests, setup instructions and a non-diagnostic listening dashboard.','Prepare an academic report with authentic citations and no invented experimental results.'])
heading('1.5 Scope and Applications')
para('The defined scope covers offline processing of WAV recordings resembling those in HLS-CMDS. The model accepts short mono audio, resamples if necessary, applies spectral transforms and returns estimated sources. Practical applications are limited to training, teaching, data analysis and methods research. General-purpose clinical separation, classification of illness and device deployment are outside the tested scope.')

chapter('2: Literature Survey')
heading('2.1 Introduction')
para('Audio separation research has developed through traditional signal decomposition, neural source models and hybrid signal-processing systems. In cardiopulmonary recordings, the task is complicated by overlapping source frequencies, source intermittency, variable recording positions and the possibility that separately recorded ground truth has phase or timing differences from the original mixture.')
heading('2.2 Review of Existing Research')
para('Tsai and collaborators proposed a periodicity-coded deep autoencoder for blind monaural heart-lung separation in 2020. Their approach motivates learning physiologically meaningful regularities rather than relying on a single fixed frequency threshold [2]. It also underscores an important conceptual distinction: a method can be designed for unpaired or blind separation, while PulseSplit uses paired targets supplied by HLS-CMDS.')
para('Ronneberger, Fischer and Brox introduced the U-Net pattern with a contracting encoder and expanding decoder joined by skip connections [3]. Although the original application was biomedical image segmentation, time-frequency spectrograms also have a two-dimensional grid on which convolutions and skip connections can operate. PulseSplit adapts this general pattern to predict two soft masks rather than a segmentation class for every image pixel.')
para('Le Roux and collaborators analyzed the limitations of several signal-to-distortion conventions and presented a scale-invariant measure, SI-SDR [4]. Separately reporting source-wise SI-SDR and the difference from an unprocessed-mixture baseline allows transparent interpretation even when absolute values differ between sources.')
para('The HLS-CMDS dataset article by Torabi, Shirani and Reilly [1] describes mixed recordings obtained from a clinical simulator and paired heart-only and lung-only recordings captured at the same stethoscope placement. Its controlled conditions are valuable for developing source-separation algorithms, but they cannot substitute for independent clinical testing.')
heading('2.3 Comparative Study')
table(['Study','Approach','Relevance / limitation'],[['Tsai et al. (2020) [2]','Periodicity-coded deep autoencoder','Specialized monaural HS–LS separation; different experimental datasets.'],['Ronneberger et al. (2015) [3]','U-Net encoder/decoder','Architectural foundation; original task was image segmentation.'],['Le Roux et al. (2019) [4]','SI-SDR metric analysis','A disciplined separation metric; does not define clinical validity.'],['Torabi et al. (2025) [1]','HLS-CMDS dataset paper','Pairs and recording conditions; limited clinical generalizability.'],['PulseSplit (proposed)','Supervised magnitude-mask U-Net','Paired simulator-source separation, fixed split and reproducible evaluation.']],'Table 2.1: Comparative Review of Relevant Research')
heading('2.4 Limitations of Existing Methods')
para('Simple band-pass filtering assumes that heart and lung occupy separable frequency bands, which is often inadequate. Blind decomposition may exploit periodicity but can suffer ambiguities in source assignments. Deep learning can overfit small or redundant datasets, and a randomly partitioned collection of audio chunks can overestimate generalization if chunks from the same original recording appear in both training and testing.')
heading('2.5 Research Gap and Proposed Contribution')
para('The proposed contribution is not a claim to invent the U-Net or the general source-separation task. It is a clearly documented, end-to-end baseline study on the permitted HLS-CMDS source, with careful pairing, leakage-aware splitting, a lightweight trainable architecture, auditable metrics and an accessible demo. The primary research question is empirical: how much does the trained system improve source estimation compared with leaving the original mixture unchanged?')

chapter('3: System Analysis and Design')
heading('3.1 Existing System')
para('Manual auscultation and conventional signal processing can visualize or filter mixed recordings but may not provide an easily reproducible two-source reconstruction workflow. Unsupervised source separation algorithms are available, yet supervised learning becomes possible when matched example outputs are supplied. A fair comparison requires a common evaluation protocol and explicit treatment of the dataset limitations.')
heading('3.2 Proposed System')
para('PulseSplit uses the UCI recording archive, extracts nested ZIP files, discovers matched M/H/L WAV triplets, partitions recording IDs and creates short waveform segments. The U-Net receives a log-compressed magnitude spectrum and predicts heart and lung masks. These masks are applied to the complex input spectrum and converted back to audio using inverse STFT. The output includes a saved checkpoint, per-recording results, optional figures and interactive playback.')
figure('system_architecture.png','Figure 3.1: PulseSplit system architecture')
heading('3.3 Functional and Non-functional Requirements')
table(['Requirement','Description'],[['FR-1','Import the original UCI HLS-CMDS archive from an approved URL or user-provided local UCI ZIP.'],['FR-2','Match mixed recordings with the correct isolated heart and lung sources.'],['FR-3','Train an STFT-mask prediction network on fixed training IDs.'],['FR-4','Track validation loss and save the best checkpoint.'],['FR-5','Evaluate previously unseen recording IDs and generate source-wise metrics.'],['NFR-1','Separate train/validation/test at recording level to mitigate leakage.'],['NFR-2','Make runs reproducible with a fixed seed and saved split manifest.'],['NFR-3','Warn that predictions are research-only and not diagnostic.']],'Table 3.1: System Requirements')
heading('3.4 Hardware and Software Requirements')
para('The intended environment is Python 3.10 or later with PyTorch, NumPy, SciPy, soundfile, scikit-learn, Matplotlib and optionally Streamlit. Training can run on a multicore CPU but may benefit from a CUDA-enabled GPU. Expected runtimes, peak memory and versions must be recorded from the final experiment environment instead of guessed.')
heading('3.5 System Architecture and Model Design')
figure('model_layers.png','Figure 3.2: Two-output spectrogram-mask U-Net architecture')
para('The encoder uses three convolutional stages with progressively increasing feature channels. A deeper bottleneck allows longer-range context. The decoder combines upsampled features with encoder activations using skip connections. A final 1-by-1 convolution predicts two channels, and a softmax across those channels generates non-negative complementary masks.')
heading('3.6 Data Flow and Evaluation Workflow')
figure('experiment_flow.png','Figure 3.3: Recording-level training, validation and test workflow')
para('The saved split manifest is authoritative: it is created before preprocessing audio into 2-second examples. After training, the test evaluation loads the recorded IDs from disk rather than generating a fresh random partition. This safeguards test integrity and enables verification of reproducibility.')

chapter('4: Dataset and Methodology')
heading('4.1 Dataset Description and Source')
para('The single permitted training data source is HLS-CMDS from the UCI Machine Learning Repository (dataset ID 1202). UCI describes 535 mono WAV recordings lasting 15 seconds at 22,050 Hz: 50 heart-only, 50 lung-only and 145 matched triplets consisting of a mixed signal and individually recorded heart and lung sources. The dataset license displayed by UCI is Creative Commons Attribution 4.0. Attribution and provenance are included in the repository [1].')
table(['Subset','Published count','Role'],[['Additional heart-only recordings','50','Reference / descriptive analysis'],['Additional lung-only recordings','50','Reference / descriptive analysis'],['Mixed recordings','145','Model inputs'],['Paired heart sources','145','Supervised targets'],['Paired lung sources','145','Supervised targets'],['Total WAV recordings','535','All released recordings']],'Table 4.1: Published Dataset Composition (UCI)')
para('The corresponding paired files use M0001.wav, H0001.wav and L0001.wav-style identification, according to the dataset paper [1]. The loader validates one-to-one matching and fails with an explicit error if a mixed recording lacks unambiguous sources. The project must not silently substitute arbitrary heart-only or lung-only files for missing supervision.')
heading('4.2 Data Collection and Preprocessing')
para('The original audio was collected using a digital stethoscope on a clinical manikin under controlled conditions; PulseSplit does not collect new medical data. Source data are downloaded from the UCI entry or accepted as a local copy of that same UCI archive. The software converts stereo signals to mono if needed, resamples all triplets to 8,000 Hz and clips three signals to a common length.')
para('Each triplet is scaled using a single factor derived from the mixture. Applying independent amplitude normalization to heart and lung signals would change the relative source strengths and make the target estimation problem inconsistent. Short contiguous segments are generated only after assigning each recording ID to one of the train, validation or test groups.')
heading('4.3 Feature Representation')
para('The short-time Fourier transform maps waveform samples to complex time-frequency coefficients. PulseSplit uses a Hann window, 512-point FFT and hop size 128 samples at 8,000 Hz. The neural network receives the log1p of the mixture magnitude. Its masks multiply the original complex mixture STFT for waveform reconstruction by inverse STFT.')
para('This approach retains the mixture phase instead of predicting separate complex spectra. Such a choice reduces model complexity, but errors caused by imperfect phase and source alignment remain an acknowledged limitation. Magnitude prediction is optimized using a log-compressed L1 loss against separately recorded source magnitude spectra.')
heading('4.4 Proposed Deep Learning Algorithm')
para('The U-Net uses convolutional blocks with batch normalization and leaky ReLU activations. The number of base channels is configurable; the default value is 16. Encoder widths are 16, 32 and 64, the bottleneck width is 128, and the decoder contracts symmetrically. Bilinear interpolation aligns spatial dimensions when odd-length STFT dimensions are encountered.')
heading('4.5 Training Strategy and Hyperparameters')
table(['Parameter','Default configuration'],[['Resampled rate','8,000 samples / second'],['Window / STFT hop','512 / 128'],['Training segment duration','2 seconds'],['Train / validation / test','70% / 15% / 15% of recording IDs'],['Model base channels','16'],['Optimizer','AdamW'],['Learning rate','0.001'],['Weight decay','0.0001'],['Batch size','8'],['Maximum epochs','30'],['Early stopping','7 validation-loss non-improving epochs'],['Primary training loss','Mean absolute error of log1p source magnitudes'],['Random seed','42']],'Table 4.2: Model and Training Configuration')
heading('4.6 Evaluation Methodology')
para('Final evaluation loads the frozen validation-selected checkpoint and predicts the two sources for each recording in the held-out test partition. Source-wise SI-SDR is computed between each predicted waveform and its corresponding isolated target. The same score is calculated when treating the original mixed recording as the estimate. The difference is reported as SI-SDR improvement, with individual-recording CSV values retained for inspection.')
para('An overall average can hide poor estimates for certain signal types, so the final report should include per-recording distributions and listening examples. Because the source recordings may not be sample-exact mixes, the final experiment must inspect timing and phase alignment; large sample-level errors may reflect both reconstruction quality and recording mismatch.')

chapter('5: Implementation')
heading('5.1 Development Environment')
para('The repository is a conventional installable Python package. It uses a pyproject.toml to describe required libraries and exposes a command-line entry point named pulsesplit. Optional dependencies provide a Streamlit demo and pytest tests. No external data are committed in the source archive.')
heading('5.2 Tools and Libraries')
table(['Tool','Role'],[['Python','General project language'],['PyTorch','Neural network training, STFT and inference'],['NumPy / SciPy','Numerical operations and sample-rate conversion'],['soundfile','Lossless WAV input and output'],['scikit-learn','Seeded recording ID partitioning'],['Matplotlib','Architecture and results visualizations'],['Streamlit','Audio upload, playback and spectrogram comparison'],['pytest','Automated software verification']],'Table 5.1: Implemented Tools and Libraries')
heading('5.3 Implementation of Individual Modules')
para('The audio module performs soundfile loading, mono conversion, polyphase resampling, matched length handling and SI-SDR. The dataset module safely extracts ZIP archives, discovers paired WAV triplets and builds recording-level splits. The model module defines the U-Net, STFT and inverse STFT. The pipeline module trains the network, tracks validation loss, saves artifacts and evaluates unseen recordings.')
para('The CLI supports download, inspect, train and evaluate commands. The Streamlit app loads a saved checkpoint and displays a mixed signal and both estimated output signals. The interface is designed for reproducible academic demonstration, not healthcare decision support.')
heading('5.4 Software Tests and Validation')
para('Three automated code tests currently pass: (1) triplet discovery and split disjointness, (2) U-Net mask shape, softmax summation, loss and reconstruction, and (3) amplitude sharing and SI-SDR behavior. A separate one-epoch pipeline smoke test ran end to end using artificial sinusoids solely to verify code execution. Those fixture values are not included as study results.')
heading('5.5 Application Interface')
para('Once a real-data checkpoint exists, running streamlit run app.py starts a local interface. The user may upload a WAV file, listen to the original waveform and the two estimated sources, and view each spectrogram. A 30-second input limit prevents unexpectedly large interactive workloads. The dashboard includes an explicit statement that it is experimental and not for medical use.')
heading('5.6 Deployment and Reproducibility')
para('The intended initial deployment is local. Saved checkpoints are not included until training with HLS-CMDS succeeds. A split.json file documents partition assignments, history.json contains epoch-level losses, and run_metadata.json stores relevant versions and counts. The results folder will contain the final per-recording CSV and summary after a genuine held-out evaluation run.')

chapter('6: Results and Discussion')
heading('6.1 Experimental Setup')
para('STATUS: PENDING REAL UCI DATASET EXECUTION. The program has passed its unit tests and a synthetic-fixture end-to-end smoke test, but the authentic HLS-CMDS dataset could not be fetched in the current execution environment. Accordingly, this draft intentionally does not report a numerical training loss, SI-SDR, source separation accuracy, qualitative superiority or comparative finding.')
heading('6.2 Training and Validation Results')
para('After running pulsesplit train on the UCI data, the final report should insert the exported epoch-loss curve and the selected best epoch. Any divergence between training and validation loss should be discussed rather than hidden. The final plot must be generated from the saved history.json file.')
heading('6.3 Evaluation Metrics and Performance Tables')
table(['Measurement','Unprocessed mixture','PulseSplit output'],[['Heart SI-SDR (dB)','PENDING','PENDING'],['Lung SI-SDR (dB)','PENDING','PENDING'],['Heart SI-SDR improvement (dB)','Not applicable','PENDING'],['Lung SI-SDR improvement (dB)','Not applicable','PENDING'],['Number of held-out recordings','PENDING SPLIT','PENDING SPLIT']],'Table 6.1: Real-Data Results to Insert After Held-Out Testing')
heading('6.4 Model Comparison and Ablation Plan')
para('At minimum, report an unprocessed-mixture baseline. Additional frequency-filtering or reduced-capacity model ablations are proposed but are not yet implemented as tested results. Any newly added baseline should be disclosed with its exact parameters and selected without referencing test-set performance.')
heading('6.5 Prediction Examples and Visualizations')
para('The completed experimental report should contain genuine examples selected using a documented protocol, such as one high-gain, one median-gain and one low-gain held-out recording. For each, show the input waveform/spectrogram, both estimated source spectrograms and corresponding target audio. No fabricated signal plots are included in this report draft.')
heading('6.6 Error Analysis and Discussion')
para('Possible failure mechanisms to investigate include overlapping heart/lung frequencies, signal-type imbalance, synchronization differences between the mixed and separately recorded sources, model capacity limits, and simulator-specific regularities. A poor result on a particular recording should be examined with listening and waveform plots rather than explained through guessed physiological causes.')

chapter('7: Conclusion and Future Scope')
heading('7.1 Summary of the Project')
para('PulseSplit now provides the engineering foundation for a reproducible deep learning source-separation study. It includes dataset acquisition from the approved UCI source, matched recording discovery, leakage-aware partitioning, spectrogram-based U-Net training, benchmark evaluation code and a local interactive demo. These components are an implementation milestone, not evidence that real-data performance objectives have already been achieved.')
heading('7.2 Achievement of Objectives')
para('Completed: core Python code, installable package, unit tests, synthetic-fixture smoke test, dataset source verification, literature grounding, report structure and draft illustrations. Pending: successful retrieval of the UCI archive, data audit, model training on the real matched audio, quality measurement, genuine result plots, and faculty-approved report details.')
heading('7.3 Limitations')
para('Only a limited controlled simulator dataset is available for the planned training and held-out testing. Individual 15-second recordings may share underlying simulator playback patterns, which could make generalization appear stronger than it is. Mask estimation reuses the mixture phase and may imperfectly reconstruct overlapping components. The proposed system has not been evaluated on actual patients or clinical environments.')
heading('7.4 Future Enhancements')
bullet(['Evaluate real-data alignment and spectral consistency before extending the model.','Experiment with stronger time-domain or complex-valued separation models only after a reliable baseline is established.','Explore training-only augmentation with gains, mild reverberation and realistic noise, while preserving label consistency.','Evaluate holdouts by sound type and anatomical position where dataset metadata supports this.','Estimate model uncertainty and identify recordings where separation confidence is low.','Validate with approved independent clinical recordings only after expert review.'])

chapter('8: References and Appendices')
heading('References (IEEE Style)')
for item in [
'[1] Y. Torabi, S. Shirani and J. P. Reilly, “Descriptor: Heart and Lung Sounds Dataset Recorded From a Clinical Manikin Using Digital Stethoscope (HLS-CMDS),” IEEE Data Descriptions, 2025, doi: 10.1109/IEEEDATA.2025.3566012. Dataset: UCI Machine Learning Repository, ID 1202, CC BY 4.0.',
'[2] K.-H. Tsai et al., “Blind Monaural Source Separation on Heart and Lung Sounds Based on Periodic-Coded Deep Autoencoder,” IEEE J. Biomed. Health Inform., vol. 24, no. 11, pp. 3203–3214, 2020, doi: 10.1109/JBHI.2020.3016831.',
'[3] O. Ronneberger, P. Fischer and T. Brox, “U-Net: Convolutional Networks for Biomedical Image Segmentation,” in MICCAI, 2015, pp. 234–241, doi: 10.1007/978-3-319-24574-4_28.',
'[4] J. Le Roux, S. Wisdom, H. Erdogan and J. R. Hershey, “SDR – Half-baked or Well Done?” in Proc. IEEE ICASSP, 2019, pp. 626–630, doi: 10.1109/ICASSP.2019.8683855.'
]:para(item)
heading('Appendix A: Key Source Code and Repository Details')
para('Repository name: pulsesplit. Primary modules: src/pulsesplit/audio.py, data.py, model.py, pipeline.py and cli.py. Interface: app.py. Installation: pip install -e ".[test,demo]". Dataset: pulsesplit download, or pulsesplit download --archive <local-UCI-archive>. Training: pulsesplit train. Evaluation: pulsesplit evaluate. Local demo: streamlit run app.py.')
heading('Appendix B: Additional Results and Screenshots')
para('RESERVED FOR AUTHENTIC EXPERIMENT OUTPUT. To be filled with loss curves, held-out metrics, spectrograms and dashboard screenshots after the real UCI dataset experiment is completed.')
heading('Appendix C: User Manual or Installation Instructions')
para('Read README.md in the repository. Install Python 3.10+, create a virtual environment, install the package, obtain the official UCI archive, inspect matching recordings, run training, evaluate the frozen checkpoint and launch the dashboard. The raw dataset must remain excluded from GitHub unless the users intentionally distribute it with full attribution and confirm license compliance.')
heading('Plagiarism/Similarity Report')
para('Placeholder for the similarity report required by the institution. No similarity score has been measured or represented in this draft.')
heading('Final Submission Evidence Checklist')
table(['Evidence item','Status'],[['Real UCI dataset obtained and provenance recorded','PENDING'],['Final dataset pairing audited','PENDING'],['Model fitted on UCI recordings','PENDING'],['Held-out predictions and SI-SDR summary exported','PENDING'],['Actual plots and dashboard screenshots added','PENDING'],['Member and guide details confirmed','PENDING'],['Faculty-approved signatures and similarity report','PENDING']])

out=OUT/'PulseSplit_Report_Working_Draft.docx';doc.save(out)
print(out)

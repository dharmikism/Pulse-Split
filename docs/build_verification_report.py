"""Build an editable software-verification appendix from one proof run.

No values are copied from an earlier run. This is deliberately a synthetic
software report; it does not calculate or claim real-data separation accuracy.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path
import xml.etree.ElementTree as ET


SOURCES = ("tests.xml", "verification.json", "history.json", "run_metadata.json", "split.json")
NOTICE = (
    "Synthetic audio software verification only. These results demonstrate execution "
    "of the implementation, not UCI dataset performance, medical accuracy, or clinical validation."
)


def _read_json(path: Path):
    def invalid_constant(value):
        raise ValueError(f"{path.name} contains a non-finite number: {value}")

    return json.loads(path.read_text(encoding="utf-8"), parse_constant=invalid_constant)


def collect_evidence(proof_dir: Path | str) -> dict:
    """Read and cross-check artifacts, failing instead of inventing missing results."""
    proof = Path(proof_dir).resolve()
    for name in SOURCES:
        if not (proof / name).is_file():
            raise ValueError(f"Missing evidence: {proof / name}")

    verification = _read_json(proof / "verification.json")
    metadata = _read_json(proof / "run_metadata.json")
    history = _read_json(proof / "history.json")
    split = _read_json(proof / "split.json")
    if "synthetic" not in str(verification.get("data", "")).lower():
        raise ValueError("This report requires explicitly labeled synthetic verification data.")
    if not isinstance(history, list) or not history:
        raise ValueError("Training history must contain at least one recorded epoch.")
    for index, row in enumerate(history, start=1):
        if row.get("epoch") != index:
            raise ValueError("Training history must contain consecutive epochs starting at 1.")
        for name in ("train_logmag_mae", "val_logmag_mae"):
            value = row.get(name)
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
                raise ValueError(f"Epoch {index} has an invalid {name}.")
    # The trainer uses a minimum improvement when replacing its checkpoint, so
    # a later, negligibly lower loss need not be the selected checkpoint's loss.
    selected_loss = math.inf
    for row in history:
        if row["val_logmag_mae"] < selected_loss - 1e-5:
            selected_loss = row["val_logmag_mae"]
    if metadata.get("best_val_logmag_mae") != selected_loss:
        raise ValueError("Selected checkpoint validation loss does not match the training selection policy.")

    split_sizes = {}
    seen = set()
    for name in ("train", "val", "test"):
        ids = split.get(name)
        if not isinstance(ids, list) or any(not isinstance(item, str) for item in ids):
            raise ValueError(f"Missing or invalid {name} recording IDs.")
        if len(set(ids)) != len(ids) or seen.intersection(ids):
            raise ValueError("Recording splits contain duplicate or overlapping IDs.")
        seen.update(ids)
        split_sizes[name] = len(ids)
    if metadata.get("recordings") != len(seen) or metadata.get("split_sizes") != split_sizes:
        raise ValueError("Recording totals do not match the recorded split IDs.")
    if metadata.get("seed") != split.get("seed"):
        raise ValueError("The metadata and recording split use different seeds.")
    shape = verification.get("inference_shape")
    if (not isinstance(shape, list) or len(shape) != 2 or shape[0] != 2
            or any(isinstance(value, bool) or not isinstance(value, int) or value <= 0 for value in shape)):
        raise ValueError("Expected inference dimensions [2 sources, positive sample count].")

    root = ET.parse(proof / "tests.xml").getroot()
    cases = list(root.iter("testcase"))
    if not cases:
        raise ValueError("JUnit evidence has no test cases.")
    tests = {"total": len(cases), "passed": 0, "failed": 0, "errors": 0, "skipped": 0}
    for case in cases:
        if case.find("error") is not None:
            tests["errors"] += 1
        elif case.find("failure") is not None:
            tests["failed"] += 1
        elif case.find("skipped") is not None:
            tests["skipped"] += 1
        else:
            tests["passed"] += 1
    # Collection errors can exist outside any testcase. Never turn those into a pass.
    suite_errors = sum(int(suite.get("errors", "0")) for suite in root.iter("testsuite"))
    tests["errors"] = max(tests["errors"], suite_errors)
    if (tests["failed"] or tests["errors"]) and verification.get("checks", {}).get("pytest") == "passed":
        raise ValueError("verification.json says pytest passed, but JUnit records failures or errors.")

    return {
        "run": proof.parent.name,
        "verification": verification,
        "metadata": metadata,
        "history": history,
        "split_sizes": split_sizes,
        "tests": tests,
        "sources": {name: hashlib.sha256((proof / name).read_bytes()).hexdigest() for name in SOURCES},
    }


def _result_rows(evidence):
    tests = evidence["tests"]
    verification = evidence["verification"]
    checks = verification.get("checks", {})
    return [
        ("Automated tests", f'{tests["passed"]} passed / {tests["total"]} collected; '
         f'{tests["failed"]} failures, {tests["errors"]} errors, {tests["skipped"]} skipped'),
        ("Completed synthetic training epochs", str(len(evidence["history"]))),
        ("Synthetic recording triplets", str(evidence["metadata"]["recordings"])),
        ("Recording split (train / validation / test)", " / ".join(
            str(evidence["split_sizes"][name]) for name in ("train", "val", "test"))),
        ("Checkpoint reload and inference", str(checks.get("checkpoint_reload_and_inference", "Not recorded"))),
        ("Inference dimensions (sources × samples)", " × ".join(map(str, verification["inference_shape"]))),
        ("Selected checkpoint synthetic validation log-magnitude MAE", str(evidence["metadata"]["best_val_logmag_mae"])),
        ("Training seed", str(evidence["metadata"]["seed"])),
        ("Training device", str(evidence["metadata"].get("device", "Not recorded"))),
    ]


def render_markdown(evidence: dict) -> str:
    """Return a portable appendix with the exact numeric values from the run."""
    verification = evidence["verification"]
    lines = ["# PulseSplit software verification report", "", f"**{NOTICE}**", "",
             f'Proof run: `{evidence["run"]}`', "",
             f'Created (UTC, recorded by the verifier): `{verification.get("created_utc", "Not recorded")}`',
             "", "## Verified software results", "", "| Measurement | Recorded result |",
             "| --- | --- |"]
    lines.extend(f"| {name} | {value} |" for name, value in _result_rows(evidence))
    lines.extend(["", "## Synthetic training loss", "",
                  "The loss is mean absolute error between log-magnitude spectra. "
                  "It is an optimization diagnostic on generated sine waves, not a separation accuracy percentage.",
                  "", "| Epoch | Training log-magnitude MAE | Validation log-magnitude MAE |",
                  "| --- | --- | --- |"])
    lines.extend(f'| {row["epoch"]} | {row["train_logmag_mae"]} | {row["val_logmag_mae"]} |'
                 for row in evidence["history"])
    lines.extend(["", "Values retain the precision present in history.json; extra digits do not imply scientific certainty.",
                  "", "## Environment", "", f'Python: `{verification.get("python", "Not recorded")}`',
                  "", f'Platform: `{verification.get("platform", "Not recorded")}`', "",
                  "| Package | Recorded version |", "| --- | --- |"])
    lines.extend(f"| {name} | {version} |" for name, version in sorted(verification.get("packages", {}).items()))
    lines.extend(["", "## Limits and pending experiments", "",
                  "The recording counts describe synthetic fixtures. The output dimensions verify that two sources "
                  "were produced, but do not show whether either source is accurate. "
                  "There are no real UCI evaluation scores in this appendix.", "",
                  "For the main experimental results, train on the documented real-data split and report held-out "
                  "heart/lung SI-SDR and improvement over the mixture baseline. Keep those results separate from "
                  "these software checks and retain the dataset provenance, model checkpoint, and evaluation outputs.",
                  "", "## Reproduce and inspect", "",
                  "Run `python tools/build_submission.py` from the project root for a new verification run. "
                  "To regenerate only this appendix from existing evidence, run "
                  '`python docs/build_verification_report.py "submission/<run>/proof"`.', "",
                  "See tests.xml for individual tests, history.json for epoch values, split.json for recording IDs, "
                  "and run_metadata.json / verification.json for settings and environment. "
                  "software_check.png, when present, is visual evidence of the synthetic execution.", "",
                  "## Source artifact hashes", "",
                  "SHA-256 values identify the exact inputs used to generate this appendix. "
                  "They detect accidental changes; they do not certify experimental quality.", "",
                  "| Source file | SHA-256 |", "| --- | --- |"])
    lines.extend(f"| {name} | `{digest}` |" for name, digest in evidence["sources"].items())
    return "\n".join(lines) + "\n"


def write_docx(evidence: dict, destination: Path) -> Path:
    """Write a Word-compatible, editable appendix; requires the reports extra."""
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError('DOCX generation needs python-docx. Install: python -m pip install -e ".[reports]"') from exc

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    for name in ("Title", "Heading 1", "Heading 2"):
        document.styles[name].font.color.rgb = RGBColor.from_string("0B5394")
    document.core_properties.title = "PulseSplit software verification report"
    document.core_properties.subject = "Synthetic execution evidence; no UCI performance claims"
    document.core_properties.comments = "Generated from the proof artifacts identified in the source hashes."

    # The Markdown and Word appendix share one rendering source so values and
    # scope labels cannot silently diverge between the two deliverables.
    lines = render_markdown(evidence).splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith("| "):
            rows = []
            while index < len(lines) and lines[index].startswith("| "):
                cells = [cell.strip().strip("`") for cell in lines[index].strip("|").split("|")]
                if not all(cell.replace("-", "").strip() == "" for cell in cells):
                    rows.append(cells)
                index += 1
            table = document.add_table(rows=1, cols=len(rows[0]))
            table.style = "Light Shading Accent 1"
            for cell, value in zip(table.rows[0].cells, rows[0]):
                cell.text = value
            for row in rows[1:]:
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = value
                    if len(value) == 64 and all(ch in "0123456789abcdef" for ch in value):
                        for run in cell.paragraphs[0].runs:
                            run.font.size = Pt(7)
            continue
        if line.startswith("# "):
            document.add_heading(line[2:], 0)
        elif line.startswith("## "):
            document.add_heading(line[3:], 1)
        elif line:
            paragraph = document.add_paragraph(line.replace("**", "").replace("`", ""))
            if line.startswith("**"):
                for run in paragraph.runs:
                    run.bold = True
        index += 1
    footer = section.footer.paragraphs[0]
    footer.text = "PulseSplit | Synthetic software verification | Editable report appendix"
    document.save(destination)
    return destination


def build_reports(proof_dir: Path | str, *, include_docx: bool = True) -> list[Path]:
    evidence = collect_evidence(proof_dir)
    proof = Path(proof_dir).resolve()
    paths = [proof / "Software_Verification_Report.md"]
    paths[0].write_text(render_markdown(evidence), encoding="utf-8")
    if include_docx:
        paths.append(write_docx(evidence, proof / "Software_Verification_Report.docx"))
    return paths


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("proof_dir", type=Path, help="One verification run's proof directory")
    parser.add_argument("--markdown-only", action="store_true", help="Generate the appendix without python-docx")
    args = parser.parse_args(argv)
    try:
        for path in build_reports(args.proof_dir, include_docx=not args.markdown_only):
            print(path)
    except (ValueError, RuntimeError, OSError, ET.ParseError) as exc:
        parser.exit(1, f"Unable to build verification report: {exc}\n")


if __name__ == "__main__":
    main()

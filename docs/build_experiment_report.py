"""Generate a checked, editable appendix for one completed dataset experiment.

This reads saved experiment outputs. It never trains a model, reruns inference,
or substitutes synthetic software verification numbers for experimental results.
"""
from __future__ import annotations

import argparse
import csv
import json
import math
from pathlib import Path

from pulsesplit.reporting import file_sha256, metric_statistics, validate_splits


METRICS = (
    ("heart_baseline_si_sdr", "Heart: mixture baseline"),
    ("heart_si_sdr", "Heart: separated estimate"),
    ("heart_improvement", "Heart: gain over mixture"),
    ("lung_baseline_si_sdr", "Lung: mixture baseline"),
    ("lung_si_sdr", "Lung: separated estimate"),
    ("lung_improvement", "Lung: gain over mixture"),
)
UCI_URL = "https://archive.ics.uci.edu/dataset/1202/hls-cmds"


def _read_json(path):
    def reject_constant(value):
        raise ValueError(f"{path.name} contains a non-finite JSON number: {value}")
    return json.loads(path.read_text(encoding="utf-8"), parse_constant=reject_constant)


def _matches(actual, expected):
    if expected is None:
        return actual is None
    if isinstance(expected, list):
        return (isinstance(actual, list) and len(actual) == len(expected)
                and all(_matches(a, b) for a, b in zip(actual, expected)))
    return (not isinstance(actual, bool) and isinstance(actual, (int, float))
            and math.isfinite(actual) and math.isclose(actual, expected, rel_tol=1e-9, abs_tol=1e-9))


def _validate_source_groups(split):
    if split.get("strategy") != "exact-source-connected-components":
        return None
    partitions = {key: name for name in ("train", "val", "test") for key in split[name]}
    groups, assignments = split.get("groups"), split.get("group_assignments")
    if not isinstance(groups, dict) or not groups or not isinstance(assignments, dict):
        raise ValueError("Exact-source split requires component membership and assignments.")
    if set(assignments) != set(groups):
        raise ValueError("Source component assignments do not cover the component IDs.")
    seen = set()
    group_counts = {name: 0 for name in ("train", "val", "test")}
    for group, members in groups.items():
        if (not isinstance(members, list) or not members or len(members) != len(set(members))
                or seen.intersection(members) or not set(members).issubset(partitions)):
            raise ValueError("Invalid or overlapping exact-source component members.")
        seen.update(members)
        assigned = assignments[group]
        if assigned not in group_counts or any(partitions[key] != assigned for key in members):
            raise ValueError("An exact-source component spans partitions or has an incorrect assignment.")
        if split.get("group_sizes", {}).get(group) != len(members):
            raise ValueError("Source component size metadata disagrees with its members.")
        group_counts[assigned] += 1
    if seen != set(partitions):
        raise ValueError("Source components do not cover all recording IDs.")
    if split.get("actual_sizes") != {name: len(split[name]) for name in group_counts}:
        raise ValueError("Source-group split size metadata disagrees with recording partitions.")
    hashes = split.get("file_hashes")
    if not isinstance(hashes, dict) or set(hashes) != set(partitions):
        raise ValueError("Source hash manifest does not cover all recording IDs.")
    owners = {}
    distinct = {role: set() for role in ("mixed", "heart", "lung")}
    for key, values in hashes.items():
        if not isinstance(values, dict) or set(values) != set(distinct):
            raise ValueError(f"Incomplete source hashes for recording {key}.")
        for role, digest in values.items():
            if (not isinstance(digest, str) or len(digest) != 64
                    or any(character not in "0123456789abcdef" for character in digest)):
                raise ValueError(f"Invalid source SHA-256 for recording {key}.")
            owner = owners.setdefault(digest, partitions[key])
            if owner != partitions[key]:
                raise ValueError("A byte-identical source hash spans experiment partitions.")
            distinct[role].add(digest)
    return {"total": len(groups), "counts": group_counts,
            "unique_files": {role: len(values) for role, values in distinct.items()}}


def collect_experiment(artifacts_dir, results_dir):
    """Reject stale hashes, partitions, and aggregate statistics before reporting."""
    artifacts, results = Path(artifacts_dir).resolve(), Path(results_dir).resolve()
    sources = {
        "artifacts/best.pt": artifacts / "best.pt",
        "artifacts/run_metadata.json": artifacts / "run_metadata.json",
        "artifacts/history.json": artifacts / "history.json",
        "artifacts/split.json": artifacts / "split.json",
        "results/summary.json": results / "summary.json",
        "results/test_per_recording.csv": results / "test_per_recording.csv",
    }
    for path in sources.values():
        if not path.is_file():
            raise ValueError(f"Missing experiment evidence: {path}")
    for name in ("dataset_provenance.json", "audio_sha256.json", "audit/dataset_audit.json", "PRETRAINING_PROTOCOL.md"):
        if (results / name).is_file():
            sources[f"results/{name}"] = results / name
    hashes = {name: file_sha256(path) for name, path in sources.items()}
    metadata = _read_json(artifacts / "run_metadata.json")
    history = _read_json(artifacts / "history.json")
    split = _read_json(artifacts / "split.json")
    summary = _read_json(results / "summary.json")
    if summary.get("checkpoint_sha256") != hashes["artifacts/best.pt"]:
        raise ValueError("Checkpoint SHA-256 differs from the evaluated checkpoint.")
    if summary.get("split_sha256") != hashes["artifacts/split.json"]:
        raise ValueError("Split SHA-256 differs from the evaluated split.")
    all_ids = [item for kind in ("train", "val", "test") for item in split.get(kind, [])]
    validate_splits(split, all_ids)
    source_groups = _validate_source_groups(split)
    split_sizes = {kind: len(split[kind]) for kind in ("train", "val", "test")}
    if metadata.get("recordings") != len(all_ids) or metadata.get("split_sizes") != split_sizes:
        raise ValueError("Training recording totals disagree with the saved split.")
    if metadata.get("seed") != split.get("seed"):
        raise ValueError("Training and split seeds disagree.")
    if ("split_strategy" in metadata and "strategy" in split
            and metadata["split_strategy"] != split["strategy"]):
        raise ValueError("Training and split strategy labels disagree.")
    audit = None
    if (results / "audit/dataset_audit.json").is_file():
        audit = _read_json(results / "audit/dataset_audit.json")
        if (audit.get("recording_count") != len(all_ids)
                or audit.get("file_count") != 3 * len(all_ids)
                or {row["id"] for row in audit.get("recordings", [])} != set(all_ids)):
            raise ValueError("Native audio audit does not match the training recording IDs/counts.")
        if audit.get("status") == "failed":
            raise ValueError("Cannot report a completed experiment from a failed native audio audit.")
        files = audit.get("files", [])
        expected_pairs = {(key, role) for key in all_ids for role in ("mixed", "heart", "lung")}
        if (len(files) != 3 * len(all_ids)
                or {(row["id"], row["role"]) for row in files} != expected_pairs):
            raise ValueError("Native audio audit file rows do not cover the experiment triplets.")
        if sorted({row["sample_rate"] for row in files}) != audit.get("sample_rates"):
            raise ValueError("Native audio audit sample-rate summary disagrees with file rows.")
        if sorted({row["channels"] for row in files}) != audit.get("channel_counts"):
            raise ValueError("Native audio audit channel summary disagrees with file rows.")
    provenance = None
    if (results / "dataset_provenance.json").is_file():
        provenance = _read_json(results / "dataset_provenance.json")
        if provenance.get("recording_triplets") != len(all_ids) or provenance.get("seed") != split.get("seed"):
            raise ValueError("Dataset provenance does not match the experiment split.")
    if not isinstance(history, list) or not history:
        raise ValueError("Training history must contain completed epochs.")
    best, selected_epoch = math.inf, None
    for index, row in enumerate(history, 1):
        if row.get("epoch") != index:
            raise ValueError("Training history has missing or reordered epochs.")
        for key in ("train_logmag_mae", "val_logmag_mae"):
            value = row.get(key)
            if (isinstance(value, bool) or not isinstance(value, (int, float))
                    or not math.isfinite(value) or value < 0):
                raise ValueError(f"Invalid training loss at epoch {index}.")
        if row["val_logmag_mae"] < best - 1e-5:
            best, selected_epoch = row["val_logmag_mae"], index
    if metadata.get("epochs_completed") != len(history) or metadata.get("best_val_logmag_mae") != best:
        raise ValueError("Training metadata disagrees with the history/checkpoint selection policy.")
    if "selected_epoch" in metadata and metadata["selected_epoch"] != selected_epoch:
        raise ValueError("Selected checkpoint epoch disagrees with the training history.")
    label = summary.get("dataset_label")
    if not isinstance(label, str) or not label.strip():
        raise ValueError("Evaluation requires a recorded dataset label.")

    with (results / "test_per_recording.csv").open(newline="", encoding="utf-8-sig") as stream:
        reader = csv.DictReader(stream)
        if not {"id", *(name for name, _ in METRICS)}.issubset(reader.fieldnames or []):
            raise ValueError("Per-recording CSV is missing required source/baseline/gain columns.")
        rows = list(reader)
    ids = [row["id"] for row in rows]
    if len(ids) != len(set(ids)) or set(ids) != set(split["test"]):
        raise ValueError("Evaluated recording IDs do not exactly match the held-out test split.")
    if summary.get("num_test_recordings") != len(rows):
        raise ValueError("Summary test count disagrees with the per-recording CSV.")
    for row in rows:
        for name, _ in METRICS:
            try:
                row[name] = float(row[name]) if row[name].strip() else math.nan
            except (ValueError, AttributeError) as exc:
                raise ValueError(f"Invalid numeric {name} for recording {row['id']}.") from exc
        for source in ("heart", "lung"):
            expected = row[f"{source}_si_sdr"] - row[f"{source}_baseline_si_sdr"]
            gain = row[f"{source}_improvement"]
            if ((math.isfinite(expected) and not _matches(gain, expected))
                    or (not math.isfinite(expected) and math.isfinite(gain))):
                raise ValueError(f"Incorrect {source} improvement for recording {row['id']}.")
    for name, _ in METRICS:
        expected = metric_statistics([row[name] for row in rows])
        saved = summary.get("statistics", {}).get(name, {})
        for statistic, value in expected.items():
            if not _matches(saved.get(statistic), value):
                raise ValueError(f"Summary {name}/{statistic} disagrees with the per-recording CSV.")
        if not _matches(summary.get("metrics", {}).get(name), expected["mean"]):
            raise ValueError(f"Summary mean for {name} disagrees with the per-recording CSV.")

    return {"metadata": metadata, "history": history, "split": split, "summary": summary,
            "rows": rows, "hashes": hashes, "selected_epoch": selected_epoch,
            "artifacts_dir": artifacts, "results_dir": results, "audit": audit, "provenance": provenance,
            "source_groups": source_groups}


def _number(value, digits=3):
    return "N/A" if value is None else f"{value:.{digits}f}"


def _escape(value):
    return str(value).replace("|", "\\|").replace("\n", " ").replace("\r", " ")


def _metric_rows(experiment):
    rows = []
    for name, label in METRICS:
        stats = experiment["summary"]["statistics"][name]
        ci = stats["mean_ci95"]
        interval = "N/A" if ci is None else f"[{ci[0]:.3f}, {ci[1]:.3f}]"
        rows.append((label, stats["count"], stats["undefined_count"], _number(stats["mean"]),
                     _number(stats["median"]), _number(stats["std"]), interval))
    return rows


def _configuration_rows(experiment):
    metadata = experiment["metadata"]
    selected = experiment["history"][experiment["selected_epoch"] - 1]
    fields = [
        ("Training data label (user supplied)", metadata.get("dataset_label", "Not recorded")),
        ("Recordings", metadata["recordings"]),
        ("Train / validation / test recording counts", " / ".join(
            str(len(experiment["split"][kind])) for kind in ("train", "val", "test"))),
        ("Split strategy", metadata.get("split_strategy", experiment["split"].get("strategy", "Recording-ID split"))),
        ("Completed epochs", len(experiment["history"])),
        ("Selected checkpoint epoch", experiment["selected_epoch"]),
        ("Selected checkpoint training log-magnitude MAE", selected["train_logmag_mae"]),
        ("Selected checkpoint validation log-magnitude MAE", selected["val_logmag_mae"]),
    ]
    groups = experiment["source_groups"]
    if groups is not None:
        fields.extend([
            ("Exact-source connected components", groups["total"]),
            ("Train / validation / test component counts", " / ".join(
                str(groups["counts"][name]) for name in ("train", "val", "test"))),
            ("Distinct byte hashes (mixture / heart / lung)", " / ".join(
                str(groups["unique_files"][role]) for role in ("mixed", "heart", "lung"))),
        ])
    for key, label in (
        ("batch_size", "Batch size"), ("base_channels", "U-Net base channels"),
        ("segment_seconds", "Training segment length (seconds)"), ("parameter_count", "Model parameters"),
        ("seed", "Seed"), ("device", "Training device"), ("runtime_seconds", "Training runtime (seconds)"),
        ("completed_utc", "Training completed (UTC)"), ("python", "Python"),
        ("platform", "Platform"), ("torch", "PyTorch"), ("numpy", "NumPy"),
    ):
        fields.append((label, metadata.get(key, "Not recorded")))
    return fields


def _split_description(experiment):
    split = experiment["split"]
    strategy = experiment["metadata"].get("split_strategy", split.get("strategy", "Recording-ID split"))
    text = "Partitions are assigned before segmentation. Recorded split strategy: " + _escape(strategy) + ". "
    if any(term in str(strategy).lower() for term in ("hash", "exact", "component")):
        text += ("Recording IDs connected through byte-identical mixed, heart, or lung files are held "
                 "in the same partition. This limits exact-source reuse across partitions. It is not a "
                 "holdout of all simulator sound types or acoustic patterns. Shared sources within a partition "
                 "also mean recording-bootstrap intervals need not have calibrated coverage for independent sources. ")
    text += ("Repeated sound types, simulator playback patterns, and anatomical locations can still span "
             "partitions. A stricter grouping study and independent patient recordings are required for broader "
             "claims. No diagnosis, treatment, or clinical validation is established by this experiment.")
    if (experiment["results_dir"] / "PRETRAINING_PROTOCOL.md").is_file():
        text += (" The pretraining protocol records a revision after duplicate-source inspection and before "
                 "training/test evaluation; see results/PRETRAINING_PROTOCOL.md for the fixed procedure.")
    return text


def _native_audio_lines(experiment):
    audit = experiment["audit"]
    if audit is None:
        return ["No native-audio audit was supplied. Native sample rate, channel count, "
                "and recording durations have not been verified by this report.", ""]
    rates = ", ".join(f"{rate:,}" for rate in audit["sample_rates"])
    durations = [row["duration_seconds"] for row in audit["files"]]
    channels = ", ".join(str(count) for count in audit["channel_counts"])
    lines = ["| Native audio measurement | Audited value |", "| --- | --- |",
             f'| Matched WAV files | {audit["file_count"]} |',
             f"| Native sample rate(s), Hz | {rates} |",
             f"| Channel count(s) | {channels} |",
             f"| File duration range, seconds | {min(durations)} to {max(durations)} |",
             f'| Total mixture duration, seconds | {audit["mixed_duration_seconds"]} |',
             f'| Audit status | {_escape(audit["status"])} |', ""]
    if audit["sample_rates"] != [22050]:
        lines.extend([
            f"The measured native rate(s), {rates} Hz, differ from the 22,050 Hz stated "
            f"in the [official UCI description]({UCI_URL}). Report the measured file properties "
            "for this downloaded release rather than substituting the catalogue rate.", "",
        ])
    if any(rate < 8000 for rate in audit["sample_rates"]):
        lines.extend(["The model pipeline resamples to 8,000 Hz. Upsampling lower-rate source audio "
                      "does not restore frequencies absent from the original recording.", ""])
    else:
        lines.extend(["The model pipeline resamples the input audio to 8,000 Hz.", ""])
    issues = audit.get("issue_counts", {})
    if issues:
        lines.extend(["Audit issue counts: " + "; ".join(f"{_escape(name)} = {count}" for name, count in sorted(issues.items())) + ". "
                      "The possible_clipping flag marks near-full-scale samples under the recorded threshold; "
                      "it is a screening warning, not proof of signal distortion.", ""])
    return lines


def render_experiment_markdown(experiment):
    summary = experiment["summary"]
    lines = ["# PulseSplit: recorded dataset experiment", "",
             f'**Evaluated data label (user supplied): {_escape(summary["dataset_label"])}**', "",
             f'Evaluation completed (UTC): {_escape(summary.get("completed_utc", "Not recorded"))}', "",
             f'Held-out test recordings: **{summary["num_test_recordings"]}**. '
             'Checkpoint and split hashes, test IDs, and aggregate statistics were cross-checked against saved artifacts.',
             "", "## Held-out SI-SDR results", "",
             "All scores and gains are in dB. Gain is separated-estimate SI-SDR minus the unprocessed-mixture "
             "baseline for the same recording and source. Positive gain means an improvement; negative gain means "
             "the estimate scores below the mixture. SI-SDR is not percentage accuracy.", "",
             "The 95% intervals below are descriptive per-recording bootstrap intervals. Shared-source "
             "dependence is not modeled; these are not calibrated intervals over independent source components.", "",
             "| Source / measurement | Valid n | Undefined | Mean (dB) | Median (dB) | SD (dB) | Mean 95% CI (dB) |",
             "| --- | --- | --- | --- | --- | --- | --- |"]
    lines.extend("| " + " | ".join(map(str, row)) + " |" for row in _metric_rows(experiment))
    lines.extend(["", "Statistics are calculated over complete held-out recordings, not training windows. "
                  "The intervals use a percentile bootstrap of recording means (2,000 resamples; seed 42). "
                  "Undefined scores are excluded and counted; intervals are unavailable with fewer than two valid recordings. "
                  "Rounded report values are backed by full-precision JSON and CSV outputs.", "",
                  "These intervals assume independent recordings. Repeated manikin playback patterns may violate "
                  "that assumption; the interval does not capture retraining variability or clinical generalization.", "",
                  "## Training and experiment configuration", "", "| Setting | Recorded value |", "| --- | --- |"])
    lines.extend(f"| {_escape(label)} | {_escape(value)} |" for label, value in _configuration_rows(experiment))
    lines.extend(["", "Checkpoint selection uses validation loss with a minimum improvement of 0.00001. "
                  "The selected checkpoint epoch can precede the final completed epoch. The training objective is "
                  "log-magnitude spectral mean absolute error; it is different from the SI-SDR evaluation metric.", "",
                  "## Completed training history", "",
                  "| Epoch | Training log-magnitude MAE | Validation log-magnitude MAE |",
                  "| --- | --- | --- |"])
    lines.extend(f'| {row["epoch"]} | {row["train_logmag_mae"]} | {row["val_logmag_mae"]} |'
                 for row in experiment["history"])
    lines.extend(["", "## Native audio audit", ""])
    lines.extend(_native_audio_lines(experiment))
    lines.extend(["## Dataset provenance and interpretation limits", "",
                  "UCI describes HLS-CMDS as digital-stethoscope recordings from a clinical manikin, "
                  "including mixed recordings and corresponding source recordings. These are manikin recordings, "
                  f"not a patient cohort. [Official UCI dataset description]({UCI_URL}).", "",
                  "The data labels in run_metadata.json and summary.json are supplied by the experiment runner. "
                  "The report verifies consistency of the saved experiment files; it does not independently "
                  "authenticate the downloaded archive or assign a clinical meaning to the scores. Retain the "
                  "download URL, archive SHA-256, audit outputs, and original dataset metadata with the submission.", "",
                  "Corresponding IDs do not by themselves prove sample-level synchronization or that the "
                  "mixture is exactly the sum of the two reference recordings. When source components are "
                  "recorded in separate acquisition phases, time/phase differences can strongly affect waveform "
                  "SI-SDR. The current pipeline resamples audio, trims to a shared length, and uses shared amplitude "
                  "scaling; it does not perform lag or phase correction. Treat scores under this protocol as "
                  "an initial benchmark, and inspect acquisition/alignment evidence before drawing source-fidelity conclusions.", "",
                  _split_description(experiment), "",
                  "Dataset attribution: Yasaman Torabi, Shahram Shirani, and James P. Reilly (2025), "
                  "HLS-CMDS; CC BY 4.0. "
                  f"[UCI attribution and license]({UCI_URL}).", "",
                  "## Held-out recording IDs", "", ", ".join(f"`{item}`" for item in experiment["split"]["test"]), "",
                  "Full per-recording results: results/test_per_recording.csv. "
                  "Full summary and intervals: results/summary.json. "
                  "Training curve: artifacts/training_history.png. Evaluation chart: results/test_comparison.png.", "",
                  "## Exact source artifacts", "",
                  "SHA-256 hashes identify the inputs used for this report. They do not certify experimental quality.", "",
                  "| Source artifact | SHA-256 |", "| --- | --- |"])
    provenance = experiment["provenance"]
    if provenance is not None:
        # Insert these before the source hash table so archive evidence is visible
        # separately from the report's own consistency checks.
        insert_at = lines.index("## Held-out recording IDs")
        lines[insert_at:insert_at] = ["## Recorded archive provenance", "",
            f'Dataset source: {_escape(provenance.get("download_url", "Not recorded"))}', "",
            f'Original archive SHA-256: `{provenance.get("archive_sha256", "Not recorded")}`', "",
            "This archive hash is preserved from the download workflow. Raw audio is not required "
            "to regenerate this report; the report does not re-download or re-authenticate it.", ""]
    lines.extend(f"| {name} | `{digest}` |" for name, digest in experiment["hashes"].items())
    lines.extend(["", "## Regenerate this appendix", "", "```powershell",
                  'python docs/build_experiment_report.py --artifacts artifacts --results results --output results/actual_experiment_report',
                  "```", "", "The Word appendix remains editable. Existing working-draft reports are not overwritten.", ""])
    return "\n".join(lines)


def _write_word(experiment, destination):
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError('Install report support with: python -m pip install -e ".[reports]"') from exc
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.65)
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10)
    for name in ("Title", "Heading 1"):
        document.styles[name].font.color.rgb = RGBColor.from_string("0B5394")
    document.core_properties.title = "PulseSplit: recorded dataset experiment"
    document.core_properties.subject = "Held-out source-separation results and experiment provenance"
    lines = render_experiment_markdown(experiment).splitlines()
    index = 0
    in_code = False
    while index < len(lines):
        line = lines[index]
        if line.startswith("| "):
            rows = []
            while index < len(lines) and lines[index].startswith("| "):
                # Escape-aware splitting keeps user-supplied labels in one cell.
                cells = [cell.strip().strip("`").replace("\\PIPE", "|") for cell in
                         lines[index].replace("\\|", "\\PIPE").strip("|").split("|")]
                if not all(cell.replace("-", "").strip() == "" for cell in cells):
                    rows.append(cells)
                index += 1
            table = document.add_table(rows=1, cols=len(rows[0]))
            table.style = "Light Shading Accent 1"
            for row_index, row in enumerate(rows):
                cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = value
                    if len(row) >= 7 or len(value) == 64:
                        for run in cell.paragraphs[0].runs:
                            run.font.size = Pt(8)
            continue
        if line.startswith("# "):
            document.add_heading(line[2:], 0)
        elif line.startswith("## "):
            document.add_heading(line[3:], 1)
        elif line.startswith("```"):
            in_code = not in_code
        elif line:
            paragraph = document.add_paragraph(line.replace("**", "").replace("`", ""))
            if line.startswith("**"):
                for run in paragraph.runs:
                    run.bold = True
            if in_code:
                for run in paragraph.runs:
                    run.font.name = "Consolas"
                    run.font.size = Pt(8)
        index += 1
    plots = ((experiment["artifacts_dir"] / "training_history.png", "Recorded training history"),
             (experiment["results_dir"] / "test_comparison.png", "Held-out source and mixture comparison"))
    present = [(path, caption) for path, caption in plots if path.is_file()]
    if present:
        document.add_page_break()
        document.add_heading("Experiment figures", 1)
        for path, caption in present:
            document.add_paragraph(caption)
            document.add_picture(str(path), width=Inches(6.7))
    section.footer.paragraphs[0].text = "PulseSplit | Recorded dataset experiment | No clinical validation"
    document.save(destination)
    return destination


def build_experiment_report(artifacts_dir, results_dir, output_dir, *, include_docx=True):
    """Return generated Markdown and (by default) editable Word report paths."""
    experiment = collect_experiment(artifacts_dir, results_dir)
    output = Path(output_dir).resolve()
    output.mkdir(parents=True, exist_ok=True)
    markdown = output / "Actual_Experiment_Report.md"
    markdown.write_text(render_experiment_markdown(experiment), encoding="utf-8")
    paths = [markdown]
    if include_docx:
        paths.append(_write_word(experiment, output / "Actual_Experiment_Report.docx"))
    return paths


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--artifacts", type=Path, default=Path("artifacts"))
    parser.add_argument("--results", type=Path, default=Path("results"))
    parser.add_argument("--output", type=Path, default=Path("results/actual_experiment_report"))
    parser.add_argument("--markdown-only", action="store_true")
    args = parser.parse_args(argv)
    try:
        for path in build_experiment_report(args.artifacts, args.results, args.output,
                                            include_docx=not args.markdown_only):
            print(path)
    except (ValueError, RuntimeError, OSError) as exc:
        parser.exit(1, f"Unable to build experiment report: {exc}\n")


if __name__ == "__main__":
    main()
